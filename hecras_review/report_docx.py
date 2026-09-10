from __future__ import annotations

from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


def _fmt(value: Any, digits: int = 3) -> str:
    if value is None:
        return "-"
    if isinstance(value, float):
        return f"{value:.{digits}f}"
    return str(value)


def _shade(cell, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_text(cell, text: Any, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(_fmt(text))
    r.bold = bold
    r.font.size = Pt(9)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _add_table(doc: Document, headers: list[str], rows: list[list[Any]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    header_row = table.rows[0]
    hdr = header_row.cells
    for i, h in enumerate(headers):
        _set_cell_text(hdr[i], h, bold=True)
        _shade(hdr[i], "D9EAF7")
    # Repeat the header row when a table flows onto another page.
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tr_pr = header_row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    for row in rows:
        data_row = table.add_row()
        cells = data_row.cells
        for i, val in enumerate(row):
            _set_cell_text(cells[i], val)
        tr_pr_data = data_row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr_data.append(cant_split)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                if i < len(row.cells):
                    row.cells[i].width = Inches(width)
    doc.add_paragraph()
    return table


def build_review_docx(report: dict[str, Any], run_history: list[dict[str, Any]] | None = None) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    for name in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
        styles[name].font.name = "Aptos Display"

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("HEC-RAS Model Review Summary")
    run.bold = True
    run.font.size = Pt(20)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Existing / Effective-like vs. Revised / Proposed Comparison").italic = True

    project = report.get("project", {})
    pair = report.get("review_pair", {})
    e = pair.get("existing", {})
    r = pair.get("revised", {})
    meta_rows = [
        ["Project", project.get("project_title") or "-"],
        ["Existing plan", f"{e.get('plan_code','')} - {e.get('plan_title','')}"],
        ["Revised plan", f"{r.get('plan_code','')} - {r.get('plan_title','')}"],
        ["Existing HEC-RAS version", e.get("program_version") or "-"],
        ["Revised HEC-RAS version", r.get("program_version") or "-"],
        ["Report generated", datetime.now().strftime("%Y-%m-%d %H:%M")],
    ]
    _add_table(doc, ["Item", "Value"], meta_rows, [2.0, 5.7])

    doc.add_heading("1. Executive Summary", level=1)
    summary = report.get("impact_summary", {})
    flags = report.get("reviewer_flags", {})
    p = doc.add_paragraph()
    p.add_run("Purpose. ").bold = True
    p.add_run(
        "This report summarizes model-input changes and hydraulic-result differences between the selected HEC-RAS plans. "
        "Reviewer flags are screening indicators only and are not an automatic FEMA, local, or other regulatory pass/fail determination."
    )
    executive_rows = [
        ["Reviewer screening flags", flags.get("flag_count", 0)],
        ["Changed 2D flow areas", len(summary.get("changed_2d_areas", []))],
        ["Modified 1D cross sections", summary.get("modified_cross_sections", 0)],
        ["Added culvert groups", summary.get("added_culvert_groups", 0)],
        ["Removed culvert groups", summary.get("removed_culvert_groups", 0)],
    ]
    _add_table(doc, ["Review metric", "Value"], executive_rows, [4.5, 1.5])

    if run_history:
        doc.add_heading("2. HEC-RAS Run History", level=1)
        rows = []
        for item in run_history:
            rows.append([
                item.get("plan_code"), item.get("status"),
                _fmt(item.get("runtime_seconds"), 1), item.get("return_code"),
                "Yes" if item.get("result_hdf_exists") else "No",
            ])
        _add_table(doc, ["Plan", "Status", "Runtime (s)", "Return code", "HDF present"], rows)
        next_section = 3
    else:
        next_section = 2

    doc.add_heading(f"{next_section}. Plan and Model Changes", level=1)
    plan_changes = report.get("plan_setting_changes", [])
    if plan_changes:
        rows = [[x.get("setting"), x.get("existing"), x.get("revised")] for x in plan_changes]
        _add_table(doc, ["Setting", "Existing", "Revised"], rows)
    else:
        doc.add_paragraph("No tracked computation-setting changes were detected between the selected plans.")

    doc.add_heading("2D Mesh Comparison", level=2)
    mesh_rows = []
    for area, m in report.get("mesh_comparison", {}).items():
        mesh_rows.append([
            area, m.get("status"), m.get("existing_cell_count"), m.get("revised_cell_count"),
            m.get("common_cell_centers"), m.get("existing_only_centers"), m.get("revised_only_centers"),
        ])
    if mesh_rows:
        _add_table(doc, ["2D area", "Status", "Existing cells", "Revised cells", "Common centers", "Existing only", "Revised only"], mesh_rows)

    doc.add_heading("1D Cross-Section Geometry", level=2)
    xs_detail = report.get("cross_section_detailed_comparison", {})
    doc.add_paragraph(
        f"Modified cross sections: {xs_detail.get('modified_count', 0)}; "
        f"added: {len(xs_detail.get('added', []))}; removed: {len(xs_detail.get('removed', []))}."
    )

    doc.add_heading("Structures and Boundary Conditions", level=2)
    structures = report.get("structure_comparison", {})
    bc = report.get("unsteady_boundary_comparison", {})
    change_rows = [
        ["Existing culvert groups", structures.get("culvert_group_count", {}).get("existing", 0)],
        ["Revised culvert groups", structures.get("culvert_group_count", {}).get("revised", 0)],
        ["Added culvert groups", len(structures.get("added_culvert_groups", []))],
        ["Modified structures", len(structures.get("modified_structures", []))],
        ["Added unsteady BC locations", len(bc.get("added_locations", []))],
        ["Removed unsteady BC locations", len(bc.get("removed_locations", []))],
        ["Existing breaklines", report.get("breaklines", {}).get("existing_count", 0)],
        ["Revised breaklines", report.get("breaklines", {}).get("revised_count", 0)],
    ]
    _add_table(doc, ["Change item", "Count"], change_rows)

    sec = next_section + 1
    doc.add_heading(f"{sec}. Hydraulic Results Summary", level=1)
    doc.add_heading("2D Results", level=2)
    h2d_rows = []
    for area, h in report.get("hydraulic_comparison", {}).get("two_d", {}).items():
        w = h.get("max_wse_delta_both_wet", {})
        wd = h.get("wet_dry_transitions", {})
        d = h.get("max_depth_delta", {})
        v = h.get("max_face_velocity_delta", {})
        ds = h.get("depth_data_source", {})
        depth_label = "Native" if not ds.get("uses_proxy") else "Derived proxy"
        h2d_rows.append([
            area, _fmt(w.get("min_delta_ft")), _fmt(w.get("max_delta_ft")),
            wd.get("became_dry_count", 0), wd.get("became_wet_count", 0),
            _fmt(d.get("min_delta_ft")), _fmt(d.get("max_delta_ft")), depth_label,
            _fmt(v.get("min_delta_ft_per_s")), _fmt(v.get("max_delta_ft_per_s")),
        ])
    if h2d_rows:
        _add_table(doc,
                   ["2D area", "Min ΔWSEL (ft)", "Max ΔWSEL (ft)", "Became dry", "Became wet", "Min ΔDepth (ft)", "Max ΔDepth (ft)", "Depth source", "Min ΔVel (ft/s)", "Max ΔVel (ft/s)"],
                   h2d_rows)

    doc.add_heading("1D Cross-Section Results", level=2)
    xsh = report.get("hydraulic_comparison", {}).get("cross_section_max_wse", {})
    xs_rows = [
        ["1D Maximum WSEL available", "Yes" if xsh.get("available", False) else "No"],
        ["Cross sections compared", xsh.get("count", 0)],
        ["Minimum ΔMaximum WSEL (ft)", _fmt(xsh.get("min_delta_ft"))],
        ["Maximum ΔMaximum WSEL (ft)", _fmt(xsh.get("max_delta_ft"))],
        ["XS with |ΔWSEL| > 0.01 ft", xsh.get("count_abs_gt_0_01_ft", 0)],
        ["XS with |ΔWSEL| > 0.1 ft", xsh.get("count_abs_gt_0_1_ft", 0)],
    ]
    if not xsh.get("available", False) and xsh.get("reason"):
        xs_rows.append(["Availability note", xsh.get("reason")])
    _add_table(doc, ["Metric", "Value"], xs_rows)
    top = xsh.get("top_10_absolute_changes", [])
    if top:
        _add_table(doc, ["River", "Reach", "RS", "Existing max WSEL", "Revised max WSEL", "ΔWSEL"], [
            [x.get("river"), x.get("reach"), x.get("rs"), _fmt(x.get("existing_max_wse_ft")), _fmt(x.get("revised_max_wse_ft")), _fmt(x.get("delta_ft"))]
            for x in top
        ])

    sec += 1
    doc.add_heading(f"{sec}. Reviewer Screening Flags", level=1)
    flag_rows = [[f.get("category"), f.get("title"), f.get("detail")] for f in flags.get("flags", [])]
    if flag_rows:
        _add_table(doc, ["Category", "Flag", "Review detail"], flag_rows)
    else:
        doc.add_paragraph("No screening flags were generated by the current ruleset.")

    sec += 1
    doc.add_heading(f"{sec}. Interpretation and Limitations", level=1)
    for item in report.get("limitations_v0_5", report.get("limitations_v0_4", [])):
        doc.add_paragraph(str(item), style="List Bullet")
    doc.add_paragraph(
        "The review tool is intended to support engineering review by organizing and visualizing model changes. "
        "Final interpretation remains the responsibility of the qualified reviewer and must follow applicable FEMA, community, agency, and project-specific criteria."
    )

    # Footer
    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run("HEC-RAS Review Tool - Screening Summary").font.size = Pt(8)

    bio = BytesIO()
    doc.save(bio)
    return bio.getvalue()


def save_review_docx(report: dict[str, Any], output_path: str | Path, run_history: list[dict[str, Any]] | None = None) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_bytes(build_review_docx(report, run_history=run_history))
    return output_path
